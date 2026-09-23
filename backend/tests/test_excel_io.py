import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.core.config import settings
from app.spreadsheet import excel_io, object_storage


def _seed_cached_formula_value(path, sheet_name: str, coordinate: str, value) -> None:
    # openpyxl itself never computes formulas, so a workbook it just created has no cached
    # value at all for a formula cell — this stands in for "a real spreadsheet application
    # already computed and saved this file at some point", which is the actual starting state
    # every real workbook this app receives is in.
    excel_io._reinject_formula_cache(path, {(sheet_name, coordinate): value})


def _build_workbook_with_cross_sheet_formula(tmp_path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1["A1"] = 5
    sheet2 = wb.create_sheet("Sheet2")
    sheet2["A1"] = "=Sheet1!A1*2"
    wb.save(path)
    wb.close()
    _seed_cached_formula_value(path, "Sheet2", "A1", 10)
    return path


def test_preserves_untouched_formula_cached_value_across_an_unrelated_edit(tmp_path):
    # Regression test for the actual bug found live: apply_edits() loads with
    # data_only=False (needed to keep formula *text* intact for editing), which never holds
    # a formula's cached *value* in the first place — an ordinary save through that path
    # silently blanked every formula cell's displayed value workbook-wide, not just whatever
    # was actually edited.
    path = _build_workbook_with_cross_sheet_formula(tmp_path)

    wb = load_workbook(path, data_only=False)
    wb["Sheet1"]["A1"].font = Font(bold=True)  # an edit that never touches the formula cell
    excel_io.save_workbook_preserving_formula_cache(wb, path, exclude=set())
    wb.close()

    values = load_workbook(path, data_only=True)
    assert values["Sheet2"]["A1"].value == 10
    assert values["Sheet1"]["A1"].value == 5

    formulas = load_workbook(path, data_only=False)
    assert formulas["Sheet2"]["A1"].value == "=Sheet1!A1*2"
    assert formulas["Sheet1"]["A1"].font.bold is True


def test_does_not_reapply_a_stale_value_to_the_cell_that_was_actually_edited(tmp_path):
    # A cell's *own* pre-edit cached value belongs to whatever it held before this edit — most
    # importantly, reapplying it would be actively wrong if the edit changed the formula
    # itself, since the old cached number no longer corresponds to the new formula at all.
    path = _build_workbook_with_cross_sheet_formula(tmp_path)

    wb = load_workbook(path, data_only=False)
    wb["Sheet2"]["A1"] = "=Sheet1!A1*3"  # a different formula than the one cached at 10
    excel_io.save_workbook_preserving_formula_cache(wb, path, exclude={("Sheet2", "A1")})
    wb.close()

    formulas = load_workbook(path, data_only=False)
    assert formulas["Sheet2"]["A1"].value == "=Sheet1!A1*3"

    values = load_workbook(path, data_only=True)
    # No formula engine recomputed this, and the stale "10" (for the *old* formula) was
    # correctly withheld — None, not a wrong number, is the honest result here.
    assert values["Sheet2"]["A1"].value is None


def test_save_survives_a_formula_cache_reinjection_failure(tmp_path, monkeypatch):
    # Regression test for a real bug: _reinject_formula_cache is a best-effort enhancement on
    # top of a structural save that has *already* succeeded and is durable on disk by the
    # time it runs — confirmed live, a real OSError there (a transient Windows sharing
    # violation) used to propagate out of save_workbook_preserving_formula_cache and abort the
    # whole apply_edits() request, even though the actual save had already landed. That left a
    # child-sheet relationship pointing at a column that had already been deleted from the
    # real file, because the caller's own post-save bookkeeping
    # (_shift_relationships_for_structural_edit) never got the chance to run.
    path = _build_workbook_with_cross_sheet_formula(tmp_path)

    wb = load_workbook(path, data_only=False)
    wb["Sheet1"]["A1"].value = 99  # an edit unrelated to the cached formula, to force a real save

    def always_fails(_path, _cached_values):
        raise PermissionError("[WinError 5] Access is denied")

    monkeypatch.setattr(excel_io, "_reinject_formula_cache", always_fails)

    excel_io.save_workbook_preserving_formula_cache(wb, path, exclude=set())  # must not raise
    wb.close()

    # The structural save itself went through fine despite the reinjection failure.
    result = load_workbook(path, data_only=False)
    assert result["Sheet1"]["A1"].value == 99


def test_workbook_write_lock_serializes_concurrent_saves_without_corrupting_the_file(tmp_path):
    # Regression test for a real bug: a production workbook was corrupted live when two
    # requests (an autosave PUT and a child-sheet creation) each ran their own
    # load-mutate-save cycle for the same file concurrently, with nothing serializing "load,
    # mutate, save" as one unit. Their writes interleaved and garbled the .xlsx's zip
    # directory structure — confirmed by hand: all of the file's individual entries were
    # still intact, only the central directory was corrupted, exactly what two overlapping
    # zipfile writers to the same path produce. workbook_write_lock() closes that window.
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.active["A1"] = 0
    wb.save(path)
    wb.close()

    errors: list[Exception] = []

    def bump():
        try:
            with excel_io.workbook_write_lock(path):
                w = excel_io.load_workbook(path, data_only=False)
                try:
                    w.active["A1"] = (w.active["A1"].value or 0) + 1
                    excel_io.save_workbook(w, path)
                finally:
                    w.close()
        except Exception as exc:  # pragma: no cover - failure path, asserted on below
            errors.append(exc)

    threads = [threading.Thread(target=bump) for _ in range(20)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert errors == []
    # The file must still be a valid, loadable workbook after 20 concurrent load-mutate-save
    # cycles — and, since workbook_write_lock() serializes them, every single increment must
    # have landed (a race would either corrupt the file outright or silently lose updates).
    result = load_workbook(path)
    assert result.active["A1"].value == 20


def test_replace_with_retry_recovers_from_a_transient_windows_sharing_violation(
    tmp_path, monkeypatch
):
    # Regression test for a real crash: a save failed outright with
    # `PermissionError: [WinError 5] Access is denied` on os.replace(), under heavy
    # concurrent read+write load on the same file — Windows (unlike POSIX) can refuse to
    # replace a file that something else briefly still has open (most plausibly antivirus
    # scanning the just-written temp file). The fix retries instead of failing immediately.
    tmp_path_file = tmp_path / "wb.xlsx.123.456.tmp"
    tmp_path_file.write_bytes(b"new content")
    real_path = tmp_path / "wb.xlsx"
    real_path.write_bytes(b"old content")

    calls: list[int] = []
    real_replace = excel_io.os.replace

    def flaky_replace(src, dst):
        calls.append(1)
        if len(calls) < 3:
            raise PermissionError("[WinError 5] Access is denied")
        return real_replace(src, dst)

    monkeypatch.setattr(excel_io.os, "replace", flaky_replace)
    monkeypatch.setattr(excel_io.time, "sleep", lambda _seconds: None)

    excel_io._replace_with_retry(tmp_path_file, real_path)

    assert len(calls) == 3
    assert real_path.read_bytes() == b"new content"


def test_replace_with_retry_reraises_once_every_attempt_is_exhausted(tmp_path, monkeypatch):
    tmp_path_file = tmp_path / "wb.xlsx.123.456.tmp"
    tmp_path_file.write_bytes(b"new content")
    real_path = tmp_path / "wb.xlsx"
    real_path.write_bytes(b"old content")

    def always_fails(src, dst):
        raise PermissionError("[WinError 5] Access is denied")

    monkeypatch.setattr(excel_io.os, "replace", always_fails)
    monkeypatch.setattr(excel_io.time, "sleep", lambda _seconds: None)

    try:
        excel_io._replace_with_retry(tmp_path_file, real_path, attempts=3)
        assert False, "expected PermissionError to propagate"
    except PermissionError:
        pass


# --- Write-path cache ---------------------------------------------------------
#
# Regression coverage for a real, measured bug: a single structural edit (insert/delete a
# row or column) correctly makes Univer recalculate every other sheet whose formulas
# reference the edited one, and each of those fires its own autosave — one user action
# cascading into several separate full-workbook load-mutate-save cycles against the same
# file, each independently paying the full reload cost. load_workbook_for_write lets the
# second and later saves in such a cascade skip their own reload by reusing the previous
# save's own in-memory Workbook object.


def test_load_workbook_for_write_reuses_the_object_from_the_immediately_preceding_save(tmp_path):
    path = tmp_path / "wb.xlsx"
    Workbook().save(path)

    first = excel_io.load_workbook_for_write(path)
    first.active["A1"] = "first save"
    excel_io.save_workbook_preserving_formula_cache(first, path)

    second = excel_io.load_workbook_for_write(path)
    # Identity, not just equal content — this is the whole point: no second parse happened.
    assert second is first
    assert second.active["A1"].value == "first save"


def test_load_workbook_for_write_reloads_fresh_once_something_else_touches_the_file(tmp_path):
    path = tmp_path / "wb.xlsx"
    Workbook().save(path)

    cached = excel_io.load_workbook_for_write(path)
    excel_io.save_workbook_preserving_formula_cache(cached, path)

    # Simulate a completely separate writer touching this same file in between — e.g. two
    # different requests racing (workbook_write_lock prevents them overlapping, but not a
    # later request from starting after an earlier one's own cache entry was written).
    other = load_workbook(path)
    other.active["A1"] = "written by someone else"
    other.save(path)
    other.close()

    reloaded = excel_io.load_workbook_for_write(path)
    assert reloaded is not cached
    assert reloaded.active["A1"].value == "written by someone else"


def test_discard_write_cache_prevents_reusing_a_workbook_left_behind_by_a_failed_save(tmp_path):
    path = tmp_path / "wb.xlsx"
    Workbook().save(path)

    failed = excel_io.load_workbook_for_write(path)
    failed.active["A1"] = "never actually saved"
    # No save_workbook_preserving_formula_cache call here — this stands in for a write that
    # raised partway through its own mutation, before ever reaching a save.
    excel_io.discard_write_cache(path)
    failed.close()

    fresh = excel_io.load_workbook_for_write(path)
    assert fresh is not failed
    # The on-disk file was never touched by the failed attempt — a fresh load reflects that,
    # not the abandoned in-memory mutation.
    assert fresh.active["A1"].value is None


# --- Remote storage mirror (STORAGE_BACKEND="s3") ---------------------------
#
# These monkeypatch object_storage's download/upload/delete with in-memory fakes rather than
# hitting a real S3-compatible endpoint — the point is confirming excel_io.py calls them at the
# right moments, not exercising boto3 itself.


def test_load_workbook_downloads_from_remote_on_local_cache_miss(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    downloads: list[str] = []

    def fake_download(key: str, local_path: Path) -> None:
        downloads.append(key)
        local_path.parent.mkdir(parents=True, exist_ok=True)
        Workbook().save(local_path)

    monkeypatch.setattr(object_storage, "download", fake_download)

    missing_path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    wb = excel_io.load_workbook(missing_path)
    wb.close()

    assert downloads == ["workbooks/owner/wb.xlsx"]


def test_load_workbook_does_not_redownload_once_present_locally(tmp_path, monkeypatch):
    # Single-process/single-instance assumption (see workbook_write_lock's docstring): once a
    # file has been fetched once in this process's lifetime, every writer re-uploads
    # immediately, so the local copy stays authoritative — no per-request network round trip.
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()

    downloads: list[str] = []
    monkeypatch.setattr(object_storage, "download", lambda key, local_path: downloads.append(key))

    wb = excel_io.load_workbook(path)
    wb.close()

    assert downloads == []


def test_save_workbook_uploads_to_remote(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    uploads: list[str] = []
    monkeypatch.setattr(object_storage, "upload", lambda local_path, key: uploads.append(key))

    path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    wb = Workbook()
    excel_io.save_workbook(wb, path)
    wb.close()

    assert uploads == ["workbooks/owner/wb.xlsx"]


def test_delete_object_deletes_from_remote(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    path.parent.mkdir(parents=True)
    path.write_bytes(b"x")

    deletes: list[str] = []
    monkeypatch.setattr(object_storage, "delete", lambda key: deletes.append(key))

    excel_io.delete_object(path)

    assert not path.exists()
    assert deletes == ["workbooks/owner/wb.xlsx"]


def test_local_backend_never_touches_object_storage(tmp_path, monkeypatch):
    # Regression guard: the default ("local") backend must be a complete no-op for every
    # object_storage function, so existing local-dev behavior is untouched by this module.
    calls: list[str] = []
    monkeypatch.setattr(object_storage, "download", lambda *a, **k: calls.append("download"))
    monkeypatch.setattr(object_storage, "upload", lambda *a, **k: calls.append("upload"))
    monkeypatch.setattr(object_storage, "delete", lambda *a, **k: calls.append("delete"))

    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    excel_io.save_workbook(wb, path)
    wb.close()
    excel_io.load_workbook(path).close()
    excel_io.delete_object(path)

    assert calls == []
