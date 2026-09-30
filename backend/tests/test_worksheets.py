import io

from fastapi.testclient import TestClient
from openpyxl import Workbook


def _upload_sample(api_client: TestClient, headers: dict, sample_xlsx_bytes: bytes) -> dict:
    response = api_client.post(
        "/workbooks",
        headers=headers,
        files={"file": ("sample.xlsx", sample_xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    return response.json()


def _get_cell(cells: list[dict], row: int, column: int) -> dict:
    return next(c for c in cells if c["row"] == row and c["column"] == column)


def test_read_worksheet_preserves_formatting_and_formulas(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    response = api_client.get(f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers)
    assert response.status_code == 200
    data = response.json()

    header_cell = _get_cell(data["cells"], row=1, column=1)
    assert header_cell["value"] == "Name"
    assert header_cell["bold"] is True
    assert header_cell["fill_color"] == "00FFFF00"

    amount_cell = _get_cell(data["cells"], row=2, column=3)
    assert amount_cell["value"] == 100
    assert amount_cell["number_format"] == "#,##0.00"

    formula_cell = _get_cell(data["cells"], row=5, column=3)
    assert formula_cell["formula"] == "=SUM(C2:C4)"
    # openpyxl never computed a result for a formula it wrote itself, so the
    # cached value is None until the file is opened in real Excel once.
    assert formula_cell["calculated_value"] is None

    assert "A5:B5" in data["merged_cells"]


def test_read_worksheet_strips_xlfn_prefix_from_formula_text(api_client: TestClient, auth_headers: dict):
    # Regression test for a real bug found in a real production workbook: a formula stored as
    # `_xlfn.IFERROR(...)` in the raw XML (Excel's own internal compatibility marker, stripped
    # by Excel itself before ever displaying the formula) reached Univer verbatim, which
    # doesn't recognize `_xlfn.IFERROR` as a function name and returned #NAME? for a formula
    # that works fine in real Excel. See cell_signal.clean_formula_text's own docstring.
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 10
    ws["A2"] = 0
    # openpyxl writes whatever string is assigned as a formula cell's value verbatim — this is
    # exactly the raw XML shape a real Excel-authored file had.
    ws["A3"] = '=+_xlfn.IFERROR(A1/A2,"")'
    buffer = io.BytesIO()
    wb.save(buffer)

    upload_response = api_client.post(
        "/workbooks",
        headers=auth_headers,
        files={
            "file": (
                "xlfn_sample.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 201
    workbook_id = upload_response.json()["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    response = api_client.get(f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers)
    assert response.status_code == 200
    formula_cell = _get_cell(response.json()["cells"], row=3, column=1)
    assert formula_cell["formula"] == '=+IFERROR(A1/A2,"")'
    assert "_xlfn" not in formula_cell["formula"]


def test_read_worksheet_strips_external_workbook_reference_index(api_client: TestClient, auth_headers: dict):
    # Regression test for a real bug found in a real production workbook (a budget document
    # assembled from several previously-separate "Sous Programme" files): a formula kept its
    # external-link syntax ('[3]Sheet!Cell') even though a sheet of that same name was copied
    # in locally, so Univer saw a reference to an unreachable external file and returned
    # #NAME?. Sets up the same shape here: a workbook with a real local sheet named "Sous
    # Programme 1" plus a formula on another sheet that references it via a bracketed index —
    # exactly what openpyxl reports for a genuine external reference — and confirms the index
    # is stripped so it resolves as an ordinary local cross-sheet reference instead.
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "Data"
    ws_main["A1"] = "='[3]Sous Programme 1'!N15"
    wb.create_sheet("Sous Programme 1")
    buffer = io.BytesIO()
    wb.save(buffer)

    upload_response = api_client.post(
        "/workbooks",
        headers=auth_headers,
        files={
            "file": (
                "external_ref_sample.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 201
    workbook_id = upload_response.json()["id"]
    worksheets = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"]
    data_worksheet_id = next(w["id"] for w in worksheets if w["name"] == "Data")

    response = api_client.get(f"/workbooks/{workbook_id}/worksheets/{data_worksheet_id}", headers=auth_headers)
    assert response.status_code == 200
    formula_cell = _get_cell(response.json()["cells"], row=1, column=1)
    assert formula_cell["formula"] == "='Sous Programme 1'!N15"
    assert "[3]" not in formula_cell["formula"]


def test_edit_worksheet_preserves_untouched_formatting(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 2, "column": 2, "value": "On Leave"}]},
    )
    assert edit_response.status_code == 200

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()

    edited_cell = _get_cell(data["cells"], row=2, column=2)
    assert edited_cell["value"] == "On Leave"

    # Untouched cells keep their formatting exactly as imported.
    header_cell = _get_cell(data["cells"], row=1, column=1)
    assert header_cell["bold"] is True
    assert header_cell["fill_color"] == "00FFFF00"
    assert "A5:B5" in data["merged_cells"]


def test_edit_worksheet_can_set_formatting_on_a_plain_cell(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={
            "edits": [
                {
                    "row": 2,
                    "column": 1,
                    "value": "Alice",
                    "bold": True,
                    "italic": True,
                    "font_color": "FFFF0000",
                    "fill_color": "FF00FF00",
                    "font_family": "Times New Roman",
                    "font_size": 14,
                    "underline": True,
                    "strikethrough": True,
                    "horizontal_alignment": "center",
                    "vertical_alignment": "top",
                    "number_format": "0.00%",
                    "borders": {"top": "thin", "bottom": None, "left": None, "right": "thick"},
                }
            ]
        },
    )
    assert edit_response.status_code == 200

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    cell = _get_cell(data["cells"], row=2, column=1)
    assert cell["value"] == "Alice"
    assert cell["bold"] is True
    assert cell["italic"] is True
    assert cell["font_color"] == "FFFF0000"
    assert cell["fill_color"] == "FF00FF00"
    assert cell["font_family"] == "Times New Roman"
    assert cell["font_size"] == 14
    assert cell["underline"] is True
    assert cell["strikethrough"] is True
    assert cell["horizontal_alignment"] == "center"
    assert cell["vertical_alignment"] == "top"
    assert cell["number_format"] == "0.00%"
    assert cell["borders"] == {"top": "thin", "bottom": None, "left": None, "right": "thick"}


def test_edit_worksheet_font_change_leaves_other_style_fields_untouched(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    # Regression test: font_family/font_size/underline/strikethrough are new, additive
    # fields on top of the existing bold/italic/color style handling — a PATCH-semantic edit
    # that only touches the new fields must not disturb bold (already set on the header row
    # by the sample fixture) or any other untouched style.
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 1, "column": 1, "font_family": "Georgia", "font_size": 16}]},
    )
    assert edit_response.status_code == 200

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    cell = _get_cell(data["cells"], row=1, column=1)
    assert cell["value"] == "Name"
    assert cell["font_family"] == "Georgia"
    assert cell["font_size"] == 16
    # Untouched — the sample fixture's header row is bold with a yellow fill.
    assert cell["bold"] is True
    assert cell["fill_color"] == "00FFFF00"


def test_edit_worksheet_value_only_edit_leaves_existing_style_untouched(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    # A CellEdit that never mentions a style field (the shape every caller sent before
    # formatting support existed) must be a pure PATCH — it should not reset the cell's
    # style to blank defaults just because those fields were omitted from the request.
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 1, "column": 1, "value": "Full Name"}]},
    )
    assert edit_response.status_code == 200

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    header_cell = _get_cell(data["cells"], row=1, column=1)
    assert header_cell["value"] == "Full Name"
    assert header_cell["bold"] is True
    assert header_cell["fill_color"] == "00FFFF00"


def test_edit_worksheet_can_explicitly_clear_bold_without_touching_color(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    # "bold": false is an explicit, meaningful value (turn bold off) — distinct from omitting
    # "bold" entirely (leave it alone). exclude_unset=True is what makes this distinguishable.
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 1, "column": 1, "value": "Name", "bold": False}]},
    )
    assert edit_response.status_code == 200

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    header_cell = _get_cell(data["cells"], row=1, column=1)
    assert header_cell["bold"] is False
    # fill_color wasn't mentioned in the edit, so it survives even though the font itself
    # was rebuilt to flip bold off.
    assert header_cell["fill_color"] == "00FFFF00"


def test_edit_worksheet_can_clear_a_cell_value(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    # Regression test: apply_cell_edits used to build edits via
    # `ws.cell(row, column, value=edit["value"])`, and openpyxl's `cell()`
    # treats `value=None` as "no value given" and silently skips the
    # assignment — so clearing a cell (typing then deleting its content)
    # looked like it saved but the old value was still on disk afterward.
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    before = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    assert _get_cell(before["cells"], row=2, column=2)["value"] == "Active"

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 2, "column": 2, "value": None}]},
    )
    assert edit_response.status_code == 200

    after = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    # The cell had no formatting of its own, so once cleared it carries no
    # signal at all and drops out of the response entirely (see
    # `_cell_has_signal`) — its absence here *is* the assertion that it
    # was actually cleared, not left at its old value.
    assert all(not (c["row"] == 2 and c["column"] == 2) for c in after["cells"])


def test_edit_worksheet_metadata_name_renames_the_sheet_and_it_survives_reload(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    # Regression test: renaming a sheet via Univer's own tab-rename UI used to not persist at
    # all — UniverSheetGrid.tsx never listened for Univer's rename mutation, so nothing was
    # ever sent to the backend and a reload silently reverted to the old name. The fix routes
    # a rename through the same generalized metadata channel merges/freeze/etc. already use
    # (WorksheetMetadataUpdate.name), not a dedicated rename endpoint.
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    edit_response = api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}",
        headers=auth_headers,
        json={"edits": [], "metadata": {"name": "Employees"}},
    )
    assert edit_response.status_code == 200
    assert edit_response.json()["name"] == "Employees"

    # Simulates a reload: a completely fresh read of both the workbook's worksheet list and
    # this worksheet's own data — neither the DB row nor the real .xlsx sheet title should
    # have reverted to the pre-rename name.
    workbook = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()
    assert workbook["worksheets"][0]["name"] == "Employees"

    data = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers
    ).json()
    assert data["name"] == "Employees"
    # The rename didn't disturb this sheet's actual content.
    assert _get_cell(data["cells"], row=1, column=1)["value"] == "Name"


def test_list_worksheet_columns_labels_by_index_not_name(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]["id"]

    response = api_client.get(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}/columns",
        headers=auth_headers,
        params={"header_start_row": 1, "header_end_row": 1},
    )
    assert response.status_code == 200
    columns = response.json()
    assert columns == [
        {"index": 1, "letter": "A", "label": "Name"},
        {"index": 2, "letter": "B", "label": "Status"},
        {"index": 3, "letter": "C", "label": "Amount"},
    ]
