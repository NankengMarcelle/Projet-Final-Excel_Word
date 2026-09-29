import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.api.routes import auth as auth_routes
from app.db.session import SessionLocal
from app.main import app
from app.models.user import User
from app.services.email_sender import ConsoleEmailSender, get_email_sender

# Tests must never depend on a real mail provider — whatever SMTP_* settings a developer's own
# .env happens to have set (e.g. for manually testing the live reset-email flow against
# Mailtrap) would otherwise make the suite make real network calls, which is slow, flaky, and
# in practice hit Mailtrap's own sandbox rate limit mid-run (confirmed live: a full test_auth.py
# run tripped "550 Too many emails per second" once real SMTP creds were in .env). This override
# keeps the suite on the fast, hermetic ConsoleEmailSender unconditionally.
app.dependency_overrides[get_email_sender] = lambda: ConsoleEmailSender()

client = TestClient(app)


@pytest.fixture
def api_client() -> TestClient:
    return client


@pytest.fixture(autouse=True)
def _reset_rate_limiters():
    # TestClient sends every request from the same client IP, so without this, register's
    # IP-only rate limit would accumulate across the whole test run (100+ tests register a
    # user) rather than being scoped per test — this isolates each test the same way a fresh
    # random email already isolates the (ip, email)-keyed login/forgot-password limiters.
    auth_routes.register_rate_limiter.reset()
    auth_routes.login_rate_limiter.reset()
    auth_routes.forgot_password_rate_limiter.reset()
    yield


def _build_sample_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    headers = ["Name", "Status", "Amount"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    rows = [("Alice", "Active", 100), ("Bob", "Inactive", 200), ("Carol", "Active", 300)]
    for row_index, (name, status, amount) in enumerate(rows, start=2):
        ws.cell(row=row_index, column=1, value=name)
        ws.cell(row=row_index, column=2, value=status)
        amount_cell = ws.cell(row=row_index, column=3, value=amount)
        amount_cell.number_format = "#,##0.00"

    ws.cell(row=5, column=2, value="Total")
    ws.cell(row=5, column=3, value="=SUM(C2:C4)")
    ws.merge_cells("A5:B5")
    ws.column_dimensions["A"].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_xlsx_bytes() -> bytes:
    return _build_sample_workbook_bytes()


def register_and_login(api_client: TestClient) -> tuple[str, str]:
    email = f"user-{uuid.uuid4().hex[:10]}@example.com"
    password = "correct-horse-battery-staple"
    api_client.post(
        "/auth/register", json={"email": email, "password": password, "full_name": "Test User"}
    )
    login_response = api_client.post("/auth/login", data={"username": email, "password": password})
    token = login_response.json()["access_token"]
    return email, token


@pytest.fixture
def auth_headers(api_client: TestClient) -> dict:
    _, token = register_and_login(api_client)
    return {"Authorization": f"Bearer {token}"}


def _promote_to_admin(email: str) -> None:
    # No app-level bootstrap path exists for creating the very first admin (that's the real-world
    # gap #10 leaves open — see admin.py's own comments) — a direct DB write here mirrors exactly
    # what a real operator has to do today to promote the first admin account.
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        user.role = "admin"
        db.commit()
    finally:
        db.close()


def register_and_login_as_admin(api_client: TestClient) -> tuple[str, str]:
    email, token = register_and_login(api_client)
    _promote_to_admin(email)
    # The access token minted at login already has this user's id as its subject — role isn't
    # baked into the JWT itself (see create_access_token), it's looked up fresh from the DB on
    # every request via get_current_user, so the existing token is valid for admin routes too
    # without needing to log in again after the promotion above.
    return email, token


@pytest.fixture
def admin_auth_headers(api_client: TestClient) -> dict:
    _, token = register_and_login_as_admin(api_client)
    return {"Authorization": f"Bearer {token}"}
