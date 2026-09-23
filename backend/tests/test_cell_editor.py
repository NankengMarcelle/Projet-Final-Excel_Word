from openpyxl import Workbook

from app.spreadsheet import cell_editor


def test_apply_cell_edits_skips_a_non_anchor_merged_cell_instead_of_crashing():
    # Regression test for a real crash found live: a plain text edit inside a large
    # multi-row/column merge produced an edit for one of the merge's own non-anchor
    # coordinates (not the anchor itself) — openpyxl's MergedCell.value is read-only, so
    # `cell.value = ...` raised AttributeError straight into a 500, turning an ordinary
    # keystroke into a permanently failing autosave (every retry hit the same crash).
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Title"
    ws.merge_cells("A1:B3")  # A1 is the anchor; A2/B1/B2/A3/B3 are non-anchor placeholders

    edits = [
        {"row": 1, "column": 1, "value": "Title changed"},  # the anchor — must still apply
        {"row": 2, "column": 2, "value": "stray edit"},  # B2 — a non-anchor placeholder
    ]

    cell_editor.apply_cell_edits(ws, edits)  # must not raise

    assert ws["A1"].value == "Title changed"


def test_apply_cell_edits_skips_style_edits_on_a_non_anchor_merged_cell_too():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Title"
    ws.merge_cells("A1:B3")

    edits = [{"row": 2, "column": 2, "bold": True, "fill_color": "FFFF0000"}]

    cell_editor.apply_cell_edits(ws, edits)  # must not raise


def test_apply_cell_edits_sets_font_family_size_underline_and_strikethrough():
    wb = Workbook()
    ws = wb.active

    cell_editor.apply_cell_edits(
        ws,
        [
            {
                "row": 1,
                "column": 1,
                "value": "Title",
                "font_family": "Georgia",
                "font_size": 18,
                "underline": True,
                "strikethrough": True,
            }
        ],
    )

    font = ws["A1"].font
    assert font.name == "Georgia"
    assert font.size == 18
    # openpyxl's own Font.underline is a style *name*, not a boolean.
    assert font.underline == "single"
    assert font.strike is True


def test_apply_cell_edits_new_font_fields_are_patch_semantic():
    # A caller that only sends font_size must not reset bold or underline already on the
    # cell — same PATCH guarantee _apply_font already made for bold/italic/font_color.
    wb = Workbook()
    ws = wb.active
    cell_editor.apply_cell_edits(
        ws, [{"row": 1, "column": 1, "value": "Title", "bold": True, "underline": True}]
    )

    cell_editor.apply_cell_edits(ws, [{"row": 1, "column": 1, "font_size": 20}])

    font = ws["A1"].font
    assert font.size == 20
    assert font.bold is True
    assert font.underline == "single"


def test_apply_cell_edits_can_turn_underline_back_off():
    wb = Workbook()
    ws = wb.active
    cell_editor.apply_cell_edits(ws, [{"row": 1, "column": 1, "value": "Title", "underline": True}])
    cell_editor.apply_cell_edits(ws, [{"row": 1, "column": 1, "underline": False}])

    assert ws["A1"].font.underline is None
