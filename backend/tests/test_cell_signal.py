from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.spreadsheet.cell_signal import cell_has_signal, clean_formula_text


def test_plain_untouched_cell_has_no_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    assert cell_has_signal(cell) is False


def test_black_font_color_is_not_signal():
    # Regression test for a real bug found in a real production workbook: a blanket
    # formatting sweep across ~500 columns wrote font_color="FF000000" (black — the default
    # text color) onto every cell it touched, with no actual data. Treating that as "real
    # content" made used_range() report the sheet's full 525-declared-column width instead of
    # its real ~18-column extent, defeating the whole point of trimming to used content.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(color="FF000000")
    assert cell_has_signal(cell) is False


def test_non_black_font_color_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(color="FFC81E1E")
    assert cell_has_signal(cell) is True


def test_bottom_vertical_alignment_is_not_signal():
    # Same category of bug as the font-color case above, on Excel's default vertical
    # alignment for an unwrapped cell.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(vertical="bottom")
    assert cell_has_signal(cell) is False


def test_top_vertical_alignment_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(vertical="top")
    assert cell_has_signal(cell) is True


def test_general_horizontal_alignment_is_not_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(horizontal="general")
    assert cell_has_signal(cell) is False


def test_center_horizontal_alignment_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(horizontal="center")
    assert cell_has_signal(cell) is True


def test_bold_cell_is_signal_even_with_default_color_and_alignment():
    # The presence of one real signal (bold) shouldn't be masked by also having
    # default-equivalent values on the other attributes.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(bold=True, color="FF000000")
    cell.alignment = Alignment(vertical="bottom")
    assert cell_has_signal(cell) is True


def test_a_value_is_always_signal_regardless_of_style():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="x")
    assert cell_has_signal(cell) is True


def test_clean_formula_text_strips_xlfn_prefix():
    # Regression test for a real bug found in a real production workbook: Excel writes some
    # functions (IFERROR among them, depending on the Excel version that saved the file) with
    # an internal `_xlfn.` compatibility prefix in the underlying XML, and silently strips it
    # again before ever displaying the formula. Nothing else did, so Univer received the raw
    # `_xlfn.IFERROR(...)` string, didn't recognize it as a function name, and returned #NAME?
    # for a formula that evaluates fine in real Excel and fine in Univer once the prefix is
    # gone — confirmed live, same cell, only that string differed.
    assert clean_formula_text('=+_xlfn.IFERROR(L8/$L$137,"")') == '=+IFERROR(L8/$L$137,"")'


def test_clean_formula_text_strips_a_nested_occurrence():
    # The prefix can sit in front of a function used as an argument deep inside a larger
    # expression, not only right after the leading "=" — a leading-prefix-only strip would
    # miss this.
    assert clean_formula_text('=SUM(A1,_xlfn.IFERROR(B1,0))') == "=SUM(A1,IFERROR(B1,0))"


def test_clean_formula_text_is_a_no_op_when_no_prefix_is_present():
    assert clean_formula_text("=SUM(A1:A10)") == "=SUM(A1:A10)"


def test_clean_formula_text_strips_external_workbook_reference_index():
    # Regression test for a real bug found in a real production workbook (a budget document
    # assembled from several previously-separate "Sous Programme" files): a formula written
    # while the target sheet still lived in a separate file kept its external-link syntax
    # ('[3]Sheet!Cell') even after a sheet of the same name was copied in locally. Confirmed
    # live with the user: real Excel shows a correct value here with zero access to the
    # external file — it's just displaying a stale cached result, not truly resolving the
    # link — and the actual local 'Sous Programme 1' tab does exist in this same workbook, so
    # stripping the index lets it resolve as a real local cross-sheet reference instead.
    assert clean_formula_text("='[3]Sous Programme 1'!N15") == "='Sous Programme 1'!N15"


def test_clean_formula_text_strips_both_xlfn_and_external_reference_together():
    assert (
        clean_formula_text('=+_xlfn.IFERROR(\'[1]Sous Programme 2\'!C9,"")')
        == '=+IFERROR(\'Sous Programme 2\'!C9,"")'
    )


def test_clean_formula_text_does_not_touch_a_normal_local_cross_sheet_reference():
    assert clean_formula_text("=SUM('Sous Programme 1'!A1:A10)") == "=SUM('Sous Programme 1'!A1:A10)"
